from __future__ import annotations

import json
import re
import shutil
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


# Rutas de entrada
PRODUCTO_ROOT = Path(
    r"C:\Users\fariass\OneDrive - SUBSECRETARIA DE SALUD PUBLICA\Escritorio\REM\REM-VIH"
)
BASE_REM_ROOT = Path(
    r"D:\DATA\REM"
)
ANIO = "2025"
SUFIJO_ANIO = ANIO[-2:]
DICCIONARIO_DIR = PRODUCTO_ROOT / ANIO
BASE_REM_DIR = BASE_REM_ROOT / f"REM_{ANIO}"
DATOS_REM_DIR = BASE_REM_DIR / "Datos"
DICT_DIR = BASE_REM_DIR / "Diccionarios"
DEIS_PATH = Path(
    r"C:\Users\fariass\OneDrive - SUBSECRETARIA DE SALUD PUBLICA\Escritorio\DATA\ESTABLECIMIENTOS\establecimientos_deis_actual.xlsx"
)

# Carpeta de salida
OUTPUT_DIR = DICCIONARIO_DIR / "salida"

# Parametros generales
REGION_RM = "13"
CHUNKSIZE = 200_000
CSV_ENCODING = "utf-8-sig"
CSV_SEP = ";"


def resolver_diccionario(pattern: str) -> Path:
    candidatos = sorted(
        DICT_DIR.glob(pattern),
        key=lambda path: (path.stat().st_mtime, path.name),
        reverse=True,
    )
    if not candidatos:
        raise FileNotFoundError(f"No se encontró un diccionario que cumpla: {pattern}")
    return candidatos[0]


DICT_A_PATH = resolver_diccionario(f"DICCIONARIO CODIGOS SA_{SUFIJO_ANIO}_*.xlsm")
DICT_P_PATH = resolver_diccionario(f"DICCIONARIO CODIGOS SP_{SUFIJO_ANIO}_*.xlsm")


REM_CONFIG = {
    "A05": {
        "csv_path": DATOS_REM_DIR / f"SerieA{ANIO}.csv",
        "dict_path": DICT_A_PATH,
        "sheet_name": "A05",
        "output_name": f"A05_{ANIO}.xlsx",
        "acumular_meses": True,
        "mes_corte": None,
    },
    "A11": {
        "csv_path": DATOS_REM_DIR / f"SerieA{ANIO}.csv",
        "dict_path": DICT_A_PATH,
        "sheet_name": "A11",
        "output_name": f"A11_{ANIO}.xlsx",
        "acumular_meses": True,
        "mes_corte": None,
    },
    "P1": {
        "csv_path": DATOS_REM_DIR / f"SerieP{ANIO}.csv",
        "dict_path": DICT_P_PATH,
        "sheet_name": "P1",
        "output_name": f"P1_{ANIO}.xlsx",
        "acumular_meses": False,
        "mes_corte": "12",
    },
    "P11": {
        "csv_path": DATOS_REM_DIR / f"SerieP{ANIO}.csv",
        "dict_path": DICT_P_PATH,
        "sheet_name": "P11",
        "output_name": f"P11_{ANIO}.xlsx",
        "acumular_meses": False,
        "mes_corte": "12",
    },
}


@dataclass(frozen=True)
class SectionTemplate:
    key: str
    name: str
    title: str
    start_row: int
    end_row: int
    first_code_row: int
    last_code_row: int
    row_to_code: dict[int, str]
    code_to_detail: dict[str, str]
    codes: list[str]
    placeholder_to_col: dict[str, int]
    columns_by_placeholder: dict[str, str]


def normalizar_espacios(texto: str) -> str:
    return re.sub(r"\s+", " ", texto).strip()


def normalizar_texto(valor: object) -> str:
    if pd.isna(valor):
        return ""
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto


def normalizar_codigo_prestacion(valor: object) -> str:
    return normalizar_texto(valor).upper()


def normalizar_texto_visible(valor: object) -> str:
    texto = normalizar_espacios(normalizar_texto(valor))
    if not texto or texto.startswith("="):
        return ""
    return texto


def es_encabezado_seccion(valor: object) -> bool:
    if valor is None:
        return False
    texto = str(valor).upper()
    return "SECCIÓN" in texto or "SECCION" in texto


def token_seccion(texto: str) -> str:
    match = re.search(r"SECCI[ÓO]N\s+([A-Z](?:\.\d+|\d*)?)", texto.upper())
    if not match:
        return "SECCION"
    return match.group(1).replace(" ", "")


def nombre_seccion(texto: str, token: str) -> str:
    limpio = normalizar_espacios(texto)
    if ":" in limpio:
        return limpio.split(":", 1)[1].strip()
    return token


def es_col_placeholder(valor: object) -> bool:
    return isinstance(valor, str) and re.fullmatch(r"COL\d{2}", valor.strip().upper()) is not None


def indice_placeholder(valor: str) -> int:
    return int(valor.strip().upper().replace("COL", ""))


def nombre_hoja(seccion: str) -> str:
    limpio = re.sub(r"[\[\]\:\*\?\/\\]", "_", str(seccion))
    return f"Seccion {limpio}"[:31]


def paths_metadata(rem: str) -> dict[str, Path]:
    prefijo = f"REM_{rem}_{ANIO}"
    return {
        "secciones": DICCIONARIO_DIR / f"{prefijo}_seccion_a_codigos.json",
        "detalle": DICCIONARIO_DIR / f"{prefijo}_codigo_a_detalle.json",
        "columnas_json": DICCIONARIO_DIR / f"{prefijo}_columnas_por_codigo.json",
        "columnas_csv": DICCIONARIO_DIR / f"{prefijo}_columnas_por_codigo.csv",
        "prestaciones_csv": DICCIONARIO_DIR / f"{prefijo}_prestaciones.csv",
        "prestaciones_json": DICCIONARIO_DIR / f"{prefijo}_prestaciones_completo.json",
    }


def cargar_deis() -> pd.DataFrame:
    columnas_deis = {
        "EstablecimientoCodigo": "Codigo DEIS",
        "EstablecimientoGlosa": "Establecimiento",
        "RegionCodigo": "Codigo Region DEIS",
        "RegionGlosa": "Región",
        "SeremiSaludCodigo_ServicioDeSaludCodigo": "Codigo Servicio DEIS",
        "SeremiSaludGlosa_ServicioDeSaludGlosa": "Servicio de Salud",
        "ComunaCodigo": "Codigo Comuna DEIS",
        "ComunaGlosa": "Comuna",
    }

    try:
        deis = pd.read_excel(
            DEIS_PATH,
            sheet_name="establecimientos",
            dtype=str,
            usecols=list(columnas_deis.keys()),
        )
    except PermissionError:
        copia_temporal = OUTPUT_DIR / "establecimientos_deis_actual_tmp.xlsx"
        copia_temporal.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(DEIS_PATH, copia_temporal)
        deis = pd.read_excel(
            copia_temporal,
            sheet_name="establecimientos",
            dtype=str,
            usecols=list(columnas_deis.keys()),
        )

    deis = deis.rename(columns=columnas_deis)

    for col in deis.columns:
        deis[col] = deis[col].map(normalizar_texto)

    deis["Codigo DEIS"] = deis["Codigo DEIS"].map(normalizar_texto)
    deis = deis.drop_duplicates(subset=["Codigo DEIS"], keep="first")
    return deis


def build_merged_lookup(ws) -> dict[tuple[int, int], object]:
    lookup: dict[tuple[int, int], object] = {}
    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        value = ws.cell(min_row, min_col).value
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                lookup[(row, col)] = value
    return lookup


def valor_visible_celda(ws, merged_lookup: dict[tuple[int, int], object], row: int, col: int) -> object:
    value = ws.cell(row, col).value
    if value not in (None, ""):
        return value
    return merged_lookup.get((row, col))


def detectar_secciones(ws, merged_lookup: dict[tuple[int, int], object]) -> list[tuple[int, str]]:
    secciones: list[tuple[int, str]] = []
    for row in range(1, ws.max_row + 1):
        texto = valor_visible_celda(ws, merged_lookup, row, 2)
        if not es_encabezado_seccion(texto):
            continue
        secciones.append((row, normalizar_espacios(str(texto))))
    return secciones


def construir_etiqueta_columna(
    ws,
    merged_lookup: dict[tuple[int, int], object],
    start_row: int,
    first_code_row: int,
    col: int,
) -> str:
    partes: list[str] = []
    for row in range(start_row + 1, first_code_row):
        texto = normalizar_texto_visible(valor_visible_celda(ws, merged_lookup, row, col))
        if not texto or es_col_placeholder(texto) or es_encabezado_seccion(texto):
            continue
        if partes and texto == partes[-1]:
            continue
        partes.append(texto)
    return " - ".join(partes)


def construir_detalle_prestacion(
    ws,
    merged_lookup: dict[tuple[int, int], object],
    row: int,
    first_placeholder_col: int,
    descriptor_context: dict[int, str],
) -> str:
    partes: list[str] = []
    for col in range(2, first_placeholder_col):
        texto = normalizar_texto_visible(valor_visible_celda(ws, merged_lookup, row, col))
        if texto:
            descriptor_context[col] = texto
        texto_contexto = descriptor_context.get(col, "")
        if not texto_contexto:
            continue
        if partes and texto_contexto == partes[-1]:
            continue
        partes.append(texto_contexto)
    return " - ".join(partes)


def extraer_plantilla_rem(rem: str, config: dict) -> tuple[list[SectionTemplate], list[dict], list[dict]]:
    wb = load_workbook(config["dict_path"], data_only=False, keep_vba=False)
    ws = wb[config["sheet_name"]]
    merged_lookup = build_merged_lookup(ws)
    secciones = detectar_secciones(ws, merged_lookup)

    plantillas: list[SectionTemplate] = []
    prestaciones_rows: list[dict] = []
    columnas_rows: list[dict] = []

    for idx, (start_row, title) in enumerate(secciones):
        hard_end = secciones[idx + 1][0] - 1 if idx + 1 < len(secciones) else ws.max_row
        key = token_seccion(title)
        section_name = nombre_seccion(title, key)

        row_to_code: dict[int, str] = {}
        placeholder_to_col: dict[str, int] = {}
        first_code_row: int | None = None
        last_code_row = start_row

        for row in range(start_row, hard_end + 1):
            code = normalizar_codigo_prestacion(ws.cell(row, 1).value)
            placeholders_en_fila: list[str] = []
            for col in range(2, ws.max_column + 1):
                valor = valor_visible_celda(ws, merged_lookup, row, col)
                if not es_col_placeholder(valor):
                    continue
                placeholder = str(valor).strip().upper()
                placeholders_en_fila.append(placeholder)
                placeholder_to_col.setdefault(placeholder, col)
            if code and placeholders_en_fila:
                row_to_code[row] = code
                if first_code_row is None:
                    first_code_row = row
                last_code_row = row

        if not row_to_code or first_code_row is None:
            continue

        columns_by_placeholder = {
            placeholder: construir_etiqueta_columna(
                ws,
                merged_lookup,
                start_row,
                first_code_row,
                col,
            )
            for placeholder, col in sorted(
                placeholder_to_col.items(),
                key=lambda item: indice_placeholder(item[0]),
            )
        }

        descriptor_context: dict[int, str] = {}
        code_to_detail: dict[str, str] = {}
        codes: list[str] = []

        for row in range(first_code_row, last_code_row + 1):
            code = row_to_code.get(row)
            if not code:
                continue

            placeholder_cols = [
                col
                for placeholder, col in placeholder_to_col.items()
                if es_col_placeholder(valor_visible_celda(ws, merged_lookup, row, col))
            ]
            if not placeholder_cols:
                continue

            first_placeholder_col = min(placeholder_cols)
            detail = construir_detalle_prestacion(
                ws,
                merged_lookup,
                row,
                first_placeholder_col,
                descriptor_context,
            )
            if not detail:
                detail = code

            code_to_detail[code] = detail
            codes.append(code)

            prestaciones_rows.append(
                {
                    "rem": rem,
                    "seccion": key,
                    "seccion_nombre": section_name,
                    "codigo_prestacion": code,
                    "detalle_prestacion": detail,
                    "detalle_niveles": detail,
                    "fila_excel": row,
                }
            )

            for placeholder, label in sorted(
                columns_by_placeholder.items(),
                key=lambda item: indice_placeholder(item[0]),
            ):
                col_idx = placeholder_to_col[placeholder]
                niveles = [parte for parte in label.split(" - ") if parte]
                columnas_rows.append(
                    {
                        "rem": rem,
                        "seccion": key,
                        "seccion_nombre": section_name,
                        "codigo_prestacion": code,
                        "detalle_prestacion": detail,
                        "codigo_columna": placeholder,
                        "columna_excel": get_column_letter(col_idx),
                        "columna_indice_excel": col_idx,
                        "columna_detalle": label,
                        "columna_categoria": niveles[0] if niveles else "",
                        "columna_subcategoria": niveles[1] if len(niveles) > 1 else "",
                        "columna_sub_subcategoria": niveles[2] if len(niveles) > 2 else "",
                        "columna_nivel_4": niveles[3] if len(niveles) > 3 else "",
                        "columna_nivel_5": niveles[4] if len(niveles) > 4 else "",
                        "columna_niveles": niveles,
                        "fila_excel": row,
                    }
                )

        plantillas.append(
            SectionTemplate(
                key=key,
                name=section_name,
                title=title,
                start_row=start_row,
                end_row=hard_end,
                first_code_row=first_code_row,
                last_code_row=last_code_row,
                row_to_code=row_to_code,
                code_to_detail=code_to_detail,
                codes=codes,
                placeholder_to_col=placeholder_to_col,
                columns_by_placeholder=columns_by_placeholder,
            )
        )

    return plantillas, prestaciones_rows, columnas_rows


def guardar_json(path: Path, data: object) -> None:
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def guardar_metadata_rem(rem: str, config: dict) -> tuple[dict[str, list[str]], dict[str, str], dict[str, dict[str, str]]]:
    plantillas, prestaciones_rows, columnas_rows = extraer_plantilla_rem(rem, config)
    paths = paths_metadata(rem)

    seccion_a_codigos = {template.key: template.codes for template in plantillas}
    codigo_a_detalle: dict[str, str] = {}
    codigo_a_seccion: dict[str, str] = {}
    columnas_por_seccion = {
        template.key: template.columns_by_placeholder for template in plantillas
    }
    columnas_por_codigo: dict[str, dict[str, str]] = {}
    columnas_completas_por_codigo: dict[str, list[dict]] = defaultdict(list)

    for template in plantillas:
        for code, detail in template.code_to_detail.items():
            codigo_a_detalle[code] = detail
            codigo_a_seccion[code] = template.key
            columnas_por_codigo[code] = template.columns_by_placeholder

    for row in columnas_rows:
        columnas_completas_por_codigo[row["codigo_prestacion"]].append(row)

    metadata_base = {
        "rem": rem,
        "anio": int(ANIO),
        "archivo_fuente": str(config["dict_path"]),
        "hoja": config["sheet_name"],
        "generado_en": datetime.now().isoformat(timespec="seconds"),
    }

    columnas_json = {
        "metadata": metadata_base,
        "columnas_por_seccion": columnas_por_seccion,
        "columnas_por_codigo": columnas_por_codigo,
        "columnas_completas_por_codigo": dict(columnas_completas_por_codigo),
    }
    prestaciones_json = {
        "metadata": metadata_base,
        "secciones": {
            template.key: {
                "nombre": template.name,
                "titulo": template.title,
                "fila_inicio": template.start_row,
                "fila_fin": template.last_code_row,
                "codigos": template.codes,
                "columnas": template.columns_by_placeholder,
            }
            for template in plantillas
        },
        "seccion_a_codigos": seccion_a_codigos,
        "codigo_a_detalle": codigo_a_detalle,
        "codigo_a_seccion": codigo_a_seccion,
        "prestaciones_por_codigo": {
            row["codigo_prestacion"]: row for row in prestaciones_rows
        },
        "prestaciones": prestaciones_rows,
    }

    guardar_json(paths["secciones"], seccion_a_codigos)
    guardar_json(paths["detalle"], codigo_a_detalle)
    guardar_json(paths["columnas_json"], columnas_json)
    guardar_json(paths["prestaciones_json"], prestaciones_json)

    pd.DataFrame(prestaciones_rows).to_csv(
        paths["prestaciones_csv"],
        index=False,
        encoding="utf-8-sig",
    )
    pd.DataFrame(columnas_rows).to_csv(
        paths["columnas_csv"],
        index=False,
        encoding="utf-8-sig",
    )

    return seccion_a_codigos, codigo_a_detalle, columnas_por_seccion


def columnas_csv_necesarias(columnas_por_seccion: dict[str, dict[str, str]]) -> list[str]:
    codigos_columna = set()
    for columnas in columnas_por_seccion.values():
        codigos_columna.update(columnas.keys())

    columnas_rem = [f"Col{codigo[-2:]}" for codigo in sorted(codigos_columna)]
    columnas_base = [
        "Mes",
        "IdServicio",
        "Ano",
        "IdEstablecimiento",
        "CodigoPrestacion",
        "IdRegion",
        "IdComuna",
    ]
    return columnas_base + columnas_rem


def leer_rem_filtrado(
    csv_path: Path,
    codigos_solicitados: set[str],
    columnas_uso: list[str],
    mes_corte: str | None = None,
) -> pd.DataFrame:
    partes = []

    for chunk in pd.read_csv(
        csv_path,
        sep=CSV_SEP,
        encoding=CSV_ENCODING,
        dtype=str,
        keep_default_na=False,
        usecols=columnas_uso,
        chunksize=CHUNKSIZE,
    ):
        chunk["IdRegion"] = chunk["IdRegion"].map(normalizar_texto)
        chunk["Ano"] = chunk["Ano"].map(normalizar_texto)
        chunk["Mes"] = chunk["Mes"].map(normalizar_texto)
        chunk["CodigoPrestacion"] = chunk["CodigoPrestacion"].map(normalizar_codigo_prestacion)

        filtro = (
            chunk["IdRegion"].eq(REGION_RM)
            & chunk["Ano"].eq(ANIO)
            & chunk["CodigoPrestacion"].isin(codigos_solicitados)
        )
        if mes_corte is not None:
            filtro = filtro & chunk["Mes"].eq(mes_corte)

        if filtro.any():
            partes.append(chunk.loc[filtro].copy())

    if not partes:
        return pd.DataFrame(columns=columnas_uso)

    return pd.concat(partes, ignore_index=True)


def preparar_base(
    df: pd.DataFrame,
    deis: pd.DataFrame,
    codigo_a_detalle: dict[str, str],
) -> pd.DataFrame:
    df = df.copy()
    df["Codigo DEIS"] = df["IdEstablecimiento"].map(normalizar_texto)
    df["Codigo Prestacion"] = df["CodigoPrestacion"].map(normalizar_codigo_prestacion)
    df["Detalle de prestación"] = df["Codigo Prestacion"].map(codigo_a_detalle).fillna("")

    df = df.merge(deis, how="left", on="Codigo DEIS")

    for col in ["Región", "Servicio de Salud", "Comuna", "Establecimiento"]:
        df[col] = df[col].fillna("")

    df["Prestación"] = df["Codigo Prestacion"]
    df["Mes"] = df["Mes"].map(normalizar_texto)
    df["Año"] = df["Ano"].map(normalizar_texto)
    df["Codigo Servicio REM"] = df["IdServicio"].map(normalizar_texto)
    df["Codigo Comuna REM"] = df["IdComuna"].map(normalizar_texto)

    return df


def construir_hoja(
    df_base: pd.DataFrame,
    seccion: str,
    codigos_seccion: list[str],
    columnas_seccion: dict[str, str],
    acumular_meses: bool,
) -> pd.DataFrame:
    codigos_set = set(codigos_seccion)
    hoja = df_base[df_base["Codigo Prestacion"].isin(codigos_set)].copy()

    columnas_valor = []
    renombrar = {}
    for codigo_col, detalle_col in columnas_seccion.items():
        col_csv = f"Col{codigo_col[-2:]}"
        if col_csv in hoja.columns:
            nombre = f"{codigo_col} - {detalle_col}" if detalle_col else codigo_col
            columnas_valor.append(col_csv)
            renombrar[col_csv] = nombre
            hoja[col_csv] = pd.to_numeric(hoja[col_csv].replace("", pd.NA), errors="coerce")

    columnas_identificacion = [
        "Región",
        "Servicio de Salud",
        "Comuna",
        "Establecimiento",
        "Codigo DEIS",
        "Prestación",
        "Detalle de prestación",
        "Año",
        "Codigo Servicio REM",
        "Codigo Comuna REM",
    ]

    if acumular_meses:
        if not hoja.empty:
            hoja = (
                hoja.groupby(columnas_identificacion, dropna=False, as_index=False)[columnas_valor]
                .sum(min_count=1)
            )
        hoja["Mes"] = "Enero-Diciembre"
    else:
        columnas_identificacion = columnas_identificacion[:7] + ["Mes"] + columnas_identificacion[7:]

    columnas_finales = columnas_identificacion + (["Mes"] if acumular_meses else []) + columnas_valor

    hoja = hoja[columnas_finales].rename(columns=renombrar)

    orden = [
        "Servicio de Salud",
        "Comuna",
        "Establecimiento",
        "Codigo DEIS",
        "Prestación",
        "Mes",
    ]
    if not hoja.empty:
        hoja = hoja.sort_values(orden, kind="stable")

    return hoja


def ajustar_formato_excel(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    ws = writer.book[sheet_name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for idx, col_name in enumerate(df.columns, start=1):
        max_len = min(max(12, len(str(col_name)) + 2), 55)
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max_len


def generar_excel(rem: str, config: dict, deis: pd.DataFrame) -> None:
    print(f"Preparando {rem}...")
    seccion_a_codigos, codigo_a_detalle, columnas_por_seccion = guardar_metadata_rem(rem, config)

    codigos_solicitados = {
        codigo
        for codigos in seccion_a_codigos.values()
        for codigo in codigos
    }
    columnas_uso = columnas_csv_necesarias(columnas_por_seccion)

    datos = leer_rem_filtrado(
        config["csv_path"],
        codigos_solicitados,
        columnas_uso,
        mes_corte=config["mes_corte"],
    )
    datos = preparar_base(datos, deis, codigo_a_detalle)

    output_path = OUTPUT_DIR / config["output_name"]
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for seccion, codigos_seccion in seccion_a_codigos.items():
            columnas_seccion = columnas_por_seccion.get(seccion, {})
            hoja = construir_hoja(
                datos,
                seccion,
                codigos_seccion,
                columnas_seccion,
                acumular_meses=config["acumular_meses"],
            )
            sheet_name = nombre_hoja(seccion)
            hoja.to_excel(writer, sheet_name=sheet_name, index=False)
            ajustar_formato_excel(writer, sheet_name, hoja)
            print(f"  Hoja {sheet_name}: {len(hoja):,} filas")

    print(f"Archivo creado: {output_path}")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    deis = cargar_deis()

    for rem, config in REM_CONFIG.items():
        generar_excel(rem, config, deis)


if __name__ == "__main__":
    main()
